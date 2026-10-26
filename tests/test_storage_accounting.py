from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from accounting import parse_money, safe_calculate
from domain import DocumentData, TableData
from exports import EXCEL_MAX_CELL_TEXT, document_xlsx, ledger_xlsx
from storage import Repository


@pytest.fixture
def repository(tmp_path):
    repo = Repository(f"sqlite:///{tmp_path / 'accounting.sqlite3'}")
    repo.initialize()
    yield repo
    repo.close()


def add_entry(repo: Repository, owner: str, *, source_key: str = "source-1", **overrides):
    values = {
        "kind": "income",
        "amount": Decimal("1234.56"),
        "currency": "RUB",
        "entry_date": date(2026, 10, 25),
        "category": "Продажи",
        "note": "",
        "source_key": source_key,
    }
    values.update(overrides)
    return repo.add_entry(owner, **values)


def test_owner_isolation_for_documents_and_entries(repository: Repository):
    data = DocumentData(title="Счёт", text="Содержимое", tables=[], warnings=[])
    document_id = repository.add_document("owner-a", "bill.pdf", data)
    entry_id, _ = add_entry(repository, "owner-a")

    assert repository.get_document("owner-b", document_id) is None
    assert repository.documents_for_search("owner-b") == []
    assert repository.delete_document("owner-b", document_id) is False
    assert repository.entries("owner-b") == []
    assert repository.void_entry("owner-b", entry_id) is False
    assert repository.summary("owner-b") == {}
    assert repository.get_document("owner-a", document_id)["data"] == data


def test_category_is_bound_literal_not_sql(repository: Repository):
    injection = "Прочее'); DROP TABLE ledger_entries; --"
    entry_id, created = add_entry(repository, "owner", category=injection)

    assert created is True
    assert repository.entries("owner")[0]["id"] == entry_id
    assert repository.entries("owner")[0]["category"] == injection
    # A second insert proves the table is still present and writes use bound ORM values.
    assert add_entry(repository, "owner", source_key="source-2")[1] is True


def test_decimal_precision_idempotency_void_and_currency_summary(repository: Repository):
    first_id, created = add_entry(repository, "owner", amount=Decimal("0.10"), source_key="telegram:1")
    duplicate_id, duplicate_created = add_entry(
        repository,
        "owner",
        amount=Decimal("999.99"),
        source_key="telegram:1",
    )
    expense_id, _ = add_entry(
        repository,
        "owner",
        kind="expense",
        amount=Decimal("0.03"),
        source_key="telegram:2",
    )
    add_entry(
        repository,
        "owner",
        amount=Decimal("10.01"),
        currency="USD",
        source_key="telegram:3",
    )

    assert (duplicate_id, duplicate_created) == (first_id, False)
    assert repository.void_entry("owner", expense_id) is True
    assert repository.void_entry("owner", expense_id) is False
    totals = repository.summary("owner")
    assert totals["RUB"] == {
        "income": Decimal("0.10"),
        "expense": Decimal("0.00"),
        "balance": Decimal("0.10"),
        "count": 1,
    }
    assert totals["USD"]["balance"] == Decimal("10.01")
    all_rows = repository.entries("owner")
    assert any(row["id"] == expense_id and row["voided"] for row in all_rows)


def test_money_parser_and_calculator_reject_ambiguous_and_unsafe_input():
    assert parse_money("1 234,50") == Decimal("1234.50")
    assert parse_money("1234.50") == Decimal("1234.50")
    with pytest.raises(ValueError):
        parse_money("1,234.50")
    with pytest.raises(ValueError):
        parse_money("12,345")  # Not silently interpreted as a thousands comma.
    with pytest.raises(ValueError, match="Деление на ноль"):
        safe_calculate("1 / 0")
    with pytest.raises(ValueError):
        safe_calculate("__import__('os').system('echo unsafe')")
    assert safe_calculate("0.1 + 0.2") == Decimal("0.3")


def test_xlsx_formula_protection_and_leading_zero_strings():
    data = DocumentData(
        title="=HYPERLINK(\"https://bad.example\")",
        text="Полный текст",
        tables=[TableData(name="Таблица", columns=["Код", "Формула"], rows=[["00123", "=2+2"]])],
        warnings=["@warning"],
    )
    workbook = load_workbook(BytesIO(document_xlsx(data)), data_only=False)
    table = workbook["Таблица"]
    assert table["A2"].value == "00123"
    assert table["A2"].data_type == "s"
    assert table["B2"].value == "'=2+2"
    assert table["B2"].data_type == "s"
    assert workbook["Документ"]["B1"].value.startswith("'")
    assert workbook["Предупреждения"]["A2"].value == "'@warning"


def test_ledger_xlsx_formula_protection_and_decimal_boundary():
    entries = [
        {
            "id": 1,
            "kind": "income",
            "amount": Decimal("10.10"),
            "currency": "RUB",
            "entry_date": date(2026, 10, 25),
            "category": "=SUM(1,1)",
            "note": "\t@cmd",
            "voided": False,
        }
    ]
    summary = {
        "RUB": {
            "income": Decimal("10.10"),
            "expense": Decimal("0.00"),
            "balance": Decimal("10.10"),
            "count": 1,
        }
    }
    workbook = load_workbook(BytesIO(ledger_xlsx(entries, summary)), data_only=False)
    ledger = workbook["Проводки"]
    assert ledger["C2"].value == pytest.approx(10.1)
    assert ledger["E2"].value == "'=SUM(1,1)"
    assert ledger["F2"].value == "'\t@cmd"


def test_xlsx_long_formula_like_text_is_not_silently_truncated():
    original = "=" + ("x" * (EXCEL_MAX_CELL_TEXT - 1))
    data = DocumentData(title="Текст", text=original, tables=[], warnings=[])

    workbook = load_workbook(BytesIO(document_xlsx(data)), data_only=False)
    assert "Длинные значения" in workbook.sheetnames
    overflow = workbook["Длинные значения"]
    chunks = [overflow.cell(row=row, column=4).value for row in range(2, overflow.max_row + 1)]
    # The first chunk carries Excel's literal-prefix apostrophe; it is not a
    # source character.  The concatenated remaining content is exact.
    restored = chunks[0][1:] + "".join(chunks[1:])
    assert restored == original
    assert "Длинные значения" in workbook["Документ"]["A4"].value


def test_entries_limit_rejects_partial_even_with_voided_rows(repository):
    first_id, _ = add_entry(repository, "owner", source_key="limit1")
    second_id, _ = add_entry(repository, "owner", source_key="limit2")
    repository.void_entry("owner", first_id)
    repository.void_entry("owner", second_id)
    assert repository.summary("owner") == {}
    with pytest.raises(ValueError, match="сузьте период"):
        repository.entries("owner", limit=1)


def test_bounded_numbers_and_large_result_formatting():
    from accounting import money_text, vat
    with pytest.raises(ValueError):
        parse_money("9" * 40)
    assert safe_calculate(" 0.1 + 0.2 ") == Decimal("0.3")
    assert money_text(Decimal("1000000000000000000000000000000000.01")).endswith(",01")
    result = vat(Decimal("1200"), Decimal("20"))
    assert result["base"] == Decimal("1000")
    assert result["included_tax"] == Decimal("200")
    assert result["added_tax"] == Decimal("240")
