from io import BytesIO

import pytest
from openpyxl import load_workbook

from domain import TableData
from exports import chart_xlsx


def test_chart_xlsx_contains_original_rows_and_chart():
    table = TableData(
        name="Расходы",
        columns=["Категория", "Сумма"],
        rows=[["Аренда", "100.25"], ["Связь", "20.00"]],
    )

    result = chart_xlsx(table, category_column=0, value_column=1)
    workbook = load_workbook(BytesIO(result))

    assert "Данные графика" in workbook.sheetnames
    assert "График" in workbook.sheetnames
    assert workbook["Данные графика"]["A2"].value == "Аренда"
    assert workbook["Данные графика"]["B2"].value == 100.25
    assert len(workbook["График"]._charts) == 1
    assert "без группировки" in workbook["Примечание"]["A1"].value
    workbook.close()


def test_chart_rejects_ambiguous_numeric_cell():
    table = TableData(name="Тест", columns=["Категория", "Сумма"], rows=[["A", "примерно 10"]])

    with pytest.raises(ValueError, match="не распознано"):
        chart_xlsx(table, 0, 1)


def test_chart_rejects_same_axis_column():
    table = TableData(name="Тест", columns=["Значение"], rows=[["10"]])

    with pytest.raises(ValueError, match="разные столбцы"):
        chart_xlsx(table, 0, 0)
