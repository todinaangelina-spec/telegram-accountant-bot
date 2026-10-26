"""End-to-end handler routing with a fully local Telegram session."""
from decimal import Decimal
from io import BytesIO

import pytest
from aiogram import Bot, Dispatcher
from aiogram.client.session.base import BaseSession
from aiogram.methods import SendDocument, SendMessage
from aiogram.types import Chat, Message, Update, User
from openpyxl import load_workbook

from ai_service import AIService
from bot import AccessMiddleware, Application
from config import Settings
from storage import Repository


class LocalSession(BaseSession):
    def __init__(self):
        super().__init__()
        self.calls = []

    async def close(self):
        pass

    async def make_request(self, bot, method, timeout=None):
        self.calls.append(method)
        assert isinstance(method, (SendMessage, SendDocument)), "Unexpected external method"
        return Message(message_id=999, date=0, chat=Chat(id=method.chat_id, type="private"),
                       text=getattr(method, "text", None))

    async def stream_content(self, url, headers=None, timeout=30, chunk_size=65536, raise_for_status=True):
        raise AssertionError("External download attempted")
        yield b""


@pytest.mark.asyncio
async def test_commands_route_and_money_persists(tmp_path):
    settings = Settings("123456:TEST_TOKEN_NOT_REAL", frozenset({123}),
                        f"sqlite:///{tmp_path / 'db.sqlite'}")
    repo = Repository(settings.database_url)
    repo.initialize()
    ai = AIService(settings)
    session = LocalSession()
    bot = Bot(settings.bot_token, session=session)
    dispatcher = Dispatcher()
    dispatcher.message.outer_middleware(AccessMiddleware(settings.allowed_user_ids))
    app = Application(settings, repo, ai)
    dispatcher.include_router(app.router)

    async def command(text, number, owner=123):
        update = Update(update_id=number, message=Message(message_id=number, date=0,
            chat=Chat(id=owner, type="private"),
            from_user=User(id=owner, is_bot=False, first_name="Fixture"), text=text))
        await dispatcher.feed_update(bot, update)

    try:
        await command("/income 0.10 RUB 2026-10-25 Fixture | fixture note", 1)
        await command("/income 0.20 RUB 2026-10-25 Fixture", 2)
        # Same message redelivery must not add money again.
        await command("/income 0.20 RUB 2026-10-25 Fixture", 2)
        await command("/expense 0.03 RUB 2026-10-25 Fixture", 3)
        assert repo.summary(123)["RUB"]["balance"] == Decimal("0.27")
        await command("/report", 4)
        assert "0,27" in session.calls[-1].text
        await command("/ledger", 5)
        workbook = load_workbook(BytesIO(session.calls[-1].document.data))
        assert workbook["Проводки"]["H1"].value == "ID записи"
        assert workbook["Проводки"].max_row == 4
        await command("/report", 6, owner=456)
        assert "Доступ закрыт" in session.calls[-1].text
        assert repo.summary(456) == {}
    finally:
        await bot.session.close()
        await ai.close()
        repo.close()


@pytest.mark.asyncio
async def test_local_text_export_and_confirmed_delete(tmp_path):
    settings = Settings("123456:TEST_TOKEN_NOT_REAL", frozenset({123}),
                        f"sqlite:///{tmp_path / 'db.sqlite'}")
    repo = Repository(settings.database_url)
    repo.initialize()
    ai = AIService(settings)
    session = LocalSession()
    bot = Bot(settings.bot_token, session=session)
    dispatcher = Dispatcher()
    dispatcher.message.outer_middleware(AccessMiddleware(settings.allowed_user_ids))
    app = Application(settings, repo, ai)
    dispatcher.include_router(app.router)

    async def command(text, number):
        await dispatcher.feed_update(bot, Update(update_id=number, message=Message(
            message_id=number, date=0, chat=Chat(id=123, type="private"),
            from_user=User(id=123, is_bot=False, first_name="Fixture"), text=text)))

    try:
        await command("/text Нечувствительная тестовая запись", 1)
        ident = repo.list_documents(123)[0]["id"]
        assert repo.get_document(123, ident)["data"].text == "Нечувствительная тестовая запись"
        await command(f"/export {ident}", 2)
        assert isinstance(session.calls[-1], SendDocument)
        await command(f"/forget {ident}", 3)
        assert repo.get_document(123, ident) is not None
        token = app.pending_delete[123][1]
        await command(f"/forget_confirm {token}", 4)
        assert repo.get_document(123, ident) is None
    finally:
        await bot.session.close()
        await ai.close()
        repo.close()
