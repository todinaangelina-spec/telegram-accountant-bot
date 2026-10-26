"""Safe XLSX exports for parsed documents and ledger reports."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accounting import parse_money

from domain import DocumentData, TableData


EXCEL_MAX_CELL_TEXT = 32_767
MAX_LEDGER_EXPORT_ENTRIES = 10_000
MAX_COLUMN_WIDTH = 50
_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=14, bold=True)
_SUBTITLE_FONT = Font(bold=True)


def _safe_text(value: object) -> str:
    """Return a literal Excel string, neutralising formula-like user values."""
    text = value if isinstance(value, str) else str(value)
    return "'" + text if text.startswith(_DANGEROUS_PREFIXES) else text


def _sheet_name(requested: str, used: set[str], fallback: str) -> str:
    """Generate a unique sheet name that fits Excel's restrictions."""
    candidate = str(requested or fallback)
    candidate = "".join(" " if ord(char) < 32 else char for char in candidate)
    candidate = candidate.translate(str.maketrans({char: " " for char in "[]:*?/\\"}))
    candidate = candidate.strip().strip("'") or fallback
    candidate = candidate[:31]
    base = candidate
    suffix_index = 2
    while candidate.casefold() in used:
        suffix = f" ({suffix_index})"
        candidate = base[: 31 - len(suffix)].rstrip() + suffix
        suffix_index += 1
    used.add(candidate.casefold())
    return candidate


def _set_width(ws: Any, column: int, width: float) -> None:
    ws.column_dimensions[get_column_letter(column)].width = max(8, min(MAX_COLUMN_WIDTH, width))


def _style_header(cells: Iterable[Any]) -> None:
    for cell in cells:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_string_cell(
    ws: Any,
    row: int,
    column: int,
    value: str,
    overflow: list[tuple[str, str]],
    location: str,
) -> None:
    """Write text or a durable reference to separately stored long text."""
    safe_value = _safe_text(value)
    # Formula protection can add an apostrophe.  Test the final XML cell value,
    # not only the source value: openpyxl otherwise truncates the last source
    # character at Excel's 32,767-character maximum.
    if len(safe_value) <= EXCEL_MAX_CELL_TEXT:
        cell = ws.cell(row=row, column=column, value=safe_value)
        cell.number_format = "@"
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        return
    overflow_number = len(overflow) + 1
    overflow.append((location, value))
    marker = f"[Полный текст на листе «Длинные значения», запись {overflow_number}]"
    cell = ws.cell(row=row, column=column, value=marker)
    cell.number_format = "@"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_overflow_sheet(workbook: Workbook, used: set[str], overflow: list[tuple[str, str]]) -> None:
    if not overflow:
        return
    ws = workbook.create_sheet(_sheet_name("Длинные значения", used, "Длинные значения"))
    ws.append(["Запись", "Источник", "Часть", "Полный текст"])
    _style_header(ws[1])
    for record_number, (location, text) in enumerate(overflow, start=1):
        # A chunk is at most Excel's actual text-cell limit; every character is
        # retained in row order rather than truncated by openpyxl/Excel.
        # Keep one spare character for a protective leading apostrophe.
        chunk_size = EXCEL_MAX_CELL_TEXT - 1
        chunks = [text[index : index + chunk_size] for index in range(0, len(text), chunk_size)]
        for part_number, chunk in enumerate(chunks, start=1):
            row = ws.max_row + 1
            ws.cell(row=row, column=1, value=record_number)
            _write_string_cell(ws, row, 2, location, [], "")
            ws.cell(row=row, column=3, value=part_number)
            _write_string_cell(ws, row, 4, chunk, [], "")
    _set_width(ws, 1, 10)
    _set_width(ws, 2, 32)
    _set_width(ws, 3, 10)
    _set_width(ws, 4, 50)
    ws.freeze_panes = "A2"


def _save_workbook(workbook: Workbook) -> bytes:
    stream = BytesIO()
    try:
        workbook.save(stream)
        return stream.getvalue()
    finally:
        workbook.close()


def _new_workbook(first_sheet_name: str) -> tuple[Workbook, Any, set[str]]:
    workbook = Workbook()
    worksheet = workbook.active
    used: set[str] = set()
    worksheet.title = _sheet_name(first_sheet_name, used, "Лист")
    return workbook, worksheet, used


def _write_document_overview(ws: Any, data: DocumentData, overflow: list[tuple[str, str]]) -> None:
    ws.cell(row=1, column=1, value="Документ").font = _TITLE_FONT
    _write_string_cell(ws, 1, 2, data.title, overflow, "Документ!B1 (заголовок)")
    ws.cell(row=3, column=1, value="Извлечённый текст").font = _SUBTITLE_FONT
    _write_string_cell(ws, 4, 1, data.text, overflow, "Документ!A4 (извлечённый текст)")
    ws.row_dimensions[4].height = 90
    _set_width(ws, 1, 50)
    _set_width(ws, 2, 50)
    ws.freeze_panes = "A3"


def _write_table_sheet(
    workbook: Workbook,
    used: set[str],
    table: TableData,
    index: int,
    overflow: list[tuple[str, str]],
) -> None:
    name = _sheet_name(table.name, used, f"Таблица {index}")
    ws = workbook.create_sheet(name)
    for column, heading in enumerate(table.columns, start=1):
        _write_string_cell(ws, 1, column, heading, overflow, f"{name}!{get_column_letter(column)}1 (заголовок)")
    _style_header(ws[1])
    for row_number, row_values in enumerate(table.rows, start=2):
        for column, value in enumerate(row_values, start=1):
            location = f"{name}!{get_column_letter(column)}{row_number}"
            # TableData requires strings; the explicit annotation protects the
            # important leading-zero invariant rather than coercing to numbers.
            _write_string_cell(ws, row_number, column, value, overflow, location)
    for column, heading in enumerate(table.columns, start=1):
        observed = max([len(heading), *[min(len(row[column - 1]), MAX_COLUMN_WIDTH) for row in table.rows]] or [8])
        _set_width(ws, column, observed + 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def document_xlsx(data: DocumentData) -> bytes:
    """Export all parsed document text, warnings and tables to a safe workbook."""
    document = DocumentData.model_validate(data)
    workbook, overview, used = _new_workbook("Документ")
    overflow: list[tuple[str, str]] = []
    generated_warnings: list[str] = []
    _write_document_overview(overview, document, overflow)
    for index, table in enumerate(document.tables, start=1):
        _write_table_sheet(workbook, used, table, index, overflow)
    if overflow:
        generated_warnings.append(
            "Некоторые значения длиннее 32 767 символов: полный текст перенесён на лист «Длинные значения»."
        )

    all_warnings = [*document.warnings, *generated_warnings]
    if all_warnings:
        warnings_sheet = workbook.create_sheet(_sheet_name("Предупреждения", used, "Предупреждения"))
        warnings_sheet.cell(row=1, column=1, value="Предупреждения")
        _style_header(warnings_sheet[1])
        for row_number, warning in enumerate(all_warnings, start=2):
            _write_string_cell(
                warnings_sheet,
                row_number,
                1,
                warning,
                overflow,
                f"Предупреждения!A{row_number}",
            )
        _set_width(warnings_sheet, 1, 50)
        warnings_sheet.freeze_panes = "A2"
    _write_overflow_sheet(workbook, used, overflow)
    return _save_workbook(workbook)


def chart_xlsx(table: TableData, category_column: int, value_column: int) -> bytes:
    """Create an Excel workbook with source rows and a bar chart.

    Column arguments are zero-based. Chart values are for visual inspection only;
    this function never aggregates or changes the accounting journal.
    """
    table = TableData.model_validate(table)
    if not 0 <= category_column < len(table.columns):
        raise ValueError("Не найден столбец категорий")
    if not 0 <= value_column < len(table.columns):
        raise ValueError("Не найден числовой столбец")
    if category_column == value_column:
        raise ValueError("Выберите разные столбцы для подписей и чисел")
    if not table.rows:
        raise ValueError("В таблице нет строк для графика")
    if len(table.rows) > 200:
        raise ValueError("Для графика допускается до 200 строк. Сначала сократите таблицу.")

    workbook, data_sheet, used = _new_workbook("Данные графика")
    for column, heading in enumerate(table.columns, start=1):
        _write_string_cell(data_sheet, 1, column, heading, [], f"Данные графика!{column}1")
    _style_header(data_sheet[1])

    for row_number, values in enumerate(table.rows, start=2):
        for column_number, raw in enumerate(values, start=1):
            if column_number - 1 == value_column:
                try:
                    number = parse_money(raw)
                except ValueError:
                    raise ValueError(
                        f"Строка {row_number - 1}: значение «{raw[:80]}» в столбце "
                        f"«{table.columns[value_column]}» не распознано как однозначное число."
                    ) from None
                cell = data_sheet.cell(row=row_number, column=column_number, value=float(number))
                cell.number_format = "0.00"
            else:
                _write_string_cell(
                    data_sheet, row_number, column_number, raw, [],
                    f"Данные графика!{column_number}:{row_number}",
                )

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = f"{table.name[:100]}: {table.columns[value_column][:60]}"
    chart.x_axis.title = table.columns[category_column][:80]
    chart.y_axis.title = table.columns[value_column][:80]
    values = Reference(data_sheet, min_col=value_column + 1, min_row=1, max_row=len(table.rows) + 1)
    categories = Reference(data_sheet, min_col=category_column + 1, min_row=2, max_row=len(table.rows) + 1)
    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 10
    chart.width = 20
    chart.legend = None

    chart_sheet = workbook.create_sheet(_sheet_name("График", used, "График"))
    chart_sheet.add_chart(chart, "A1")
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = data_sheet.dimensions
    note = workbook.create_sheet(_sheet_name("Примечание", used, "Примечание"))
    note.append(["График показывает строки таблицы без группировки и без добавления записей в журнал."])
    note.append(["Числовой столбец преобразован в числовой формат Excel только для построения диаграммы."])
    note.append(["Сверьте исходные значения, единицы измерения, валюту и наличие строк «Итого»."])
    _set_width(note, 1, 100)
    return _save_workbook(workbook)


def _excel_number(value: object, field: str) -> float:
    """Convert only at the XLSX boundary; never use this in persistence/math."""
    if isinstance(value, bool) or isinstance(value, float):
        raise ValueError(f"{field} должен быть Decimal для экспорта")
    try:
        decimal_value = value if isinstance(value, Decimal) else Decimal(value)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValueError(f"Некорректное числовое поле {field}") from exc
    if not decimal_value.is_finite():
        raise ValueError(f"Некорректное числовое поле {field}")
    return float(decimal_value)


def _entry_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError("Дата операции должна быть date")


def ledger_xlsx(entries: list[dict[str, Any]], summary: dict[str, dict[str, Any]]) -> bytes:
    """Export a bounded ledger and exact source summary as a styled workbook.

    More than ``MAX_LEDGER_EXPORT_ENTRIES`` rows raises an explicit error rather
    than silently exporting a partial ledger whose totals could be misleading.
    """
    if not isinstance(entries, list):
        raise ValueError("entries должен быть списком")
    if len(entries) > MAX_LEDGER_EXPORT_ENTRIES:
        raise ValueError(
            f"Экспорт ограничен {MAX_LEDGER_EXPORT_ENTRIES} операциями; сузьте период перед выгрузкой"
        )
    if not isinstance(summary, dict):
        raise ValueError("summary должен быть словарём")
    summary_count = 0
    for values in summary.values():
        if not isinstance(values, dict):
            raise ValueError("Каждый итог по валюте должен быть словарём")
        count = values.get("count")
        if isinstance(count, bool) or not isinstance(count, int) or count < 0:
            raise ValueError("count должен быть неотрицательным целым числом")
        summary_count += count
    # Repository.entries includes voided rows while Repository.summary omits
    # them, hence a complete supplied ledger can legitimately be longer.  The
    # converse would mean workbook totals contain records not shown in it.
    if summary_count > len(entries):
        raise ValueError("Итоги содержат операции, отсутствующие в выгружаемом журнале")
    workbook, ledger, used = _new_workbook("Проводки")
    headings = ["Дата", "Тип", "Сумма", "Валюта", "Категория", "Примечание", "Аннулирована", "ID записи"]
    for column, heading in enumerate(headings, start=1):
        ledger.cell(row=1, column=column, value=heading)
    _style_header(ledger[1])
    for row_number, entry in enumerate(entries, start=2):
        if not isinstance(entry, dict):
            raise ValueError("Каждая операция должна быть словарём")
        ledger.cell(row=row_number, column=1, value=_entry_date(entry.get("entry_date")))
        ledger.cell(row=row_number, column=1).number_format = "DD.MM.YYYY"
        _write_string_cell(ledger, row_number, 2, str(entry.get("kind", "")), [], "")
        ledger.cell(row=row_number, column=3, value=_excel_number(entry.get("amount"), "amount"))
        ledger.cell(row=row_number, column=3).number_format = "0.00"
        _write_string_cell(ledger, row_number, 4, str(entry.get("currency", "")), [], "")
        _write_string_cell(ledger, row_number, 5, str(entry.get("category", "")), [], "")
        _write_string_cell(ledger, row_number, 6, str(entry.get("note", "")), [], "")
        _write_string_cell(ledger, row_number, 7, "Да" if entry.get("voided") else "Нет", [], "")
        ledger.cell(row=row_number, column=8, value=entry.get("id"))
    for column, width in enumerate((14, 14, 16, 12, 28, 50, 18, 14), start=1):
        _set_width(ledger, column, width)
    ledger.freeze_panes = "A2"
    ledger.auto_filter.ref = ledger.dimensions

    totals = workbook.create_sheet(_sheet_name("Итоги", used, "Итоги"))
    total_headings = ["Валюта", "Доход", "Расход", "Баланс", "Количество операций"]
    for column, heading in enumerate(total_headings, start=1):
        totals.cell(row=1, column=column, value=heading)
    _style_header(totals[1])
    for row_number, (currency, values) in enumerate(sorted(summary.items()), start=2):
        _write_string_cell(totals, row_number, 1, str(currency), [], "")
        for column, key in enumerate(("income", "expense", "balance"), start=2):
            totals.cell(row=row_number, column=column, value=_excel_number(values.get(key), key))
            totals.cell(row=row_number, column=column).number_format = "0.00"
        count = values.get("count")
        totals.cell(row=row_number, column=5, value=count)
    for column, width in enumerate((12, 18, 18, 18, 24), start=1):
        _set_width(totals, column, width)
    totals.freeze_panes = "A2"
    totals.auto_filter.ref = totals.dimensions
    return _save_workbook(workbook)
